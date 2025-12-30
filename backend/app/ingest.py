import io
import os
import re
import tempfile
from datetime import date, datetime
from typing import Optional, Tuple, List

from fastapi import APIRouter, Depends, UploadFile, File
from sqlalchemy.orm import Session

from .db import get_db, SessionLocal
from .models import Document
from .auth import get_current_user

# Optional parsers (best effort)
try:
    from docx import Document as DocxDocument
except Exception:
    DocxDocument = None

try:
    import openpyxl
except Exception:
    openpyxl = None

try:
    from pptx import Presentation
except Exception:
    Presentation = None

try:
    import extract_msg
except Exception:
    extract_msg = None

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

# EML parsing (stdlib)
from email import policy
from email.parser import BytesParser
from email.message import Message

router = APIRouter(prefix="/api/ingest", tags=["ingest"])


def _safe_print(s: str) -> None:
    try:
        print(s, flush=True)
    except Exception:
        pass


def _guess_doc_date_from_filename(name: str) -> Optional[date]:
    base = os.path.basename(name or "")

    # YYYY-MM-DD
    m = re.search(r"(20\d{2})[-_.](\d{2})[-_.](\d{2})", base)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d)
        except Exception:
            return None

    # DD.MM.YYYY or DD-MM-YYYY
    m = re.search(r"(\d{2})[.\-_/](\d{2})[.\-_/](20\d{2})", base)
    if m:
        d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d)
        except Exception:
            return None

    return None


def _html_to_text(html: str) -> str:
    if not html:
        return ""
    # remove scripts/styles
    html = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", html)
    # remove comments
    html = re.sub(r"(?is)<!--.*?-->", " ", html)
    # remove tags
    text = re.sub(r"(?s)<[^>]+>", " ", html)
    # unescape common entities
    text = (
        text.replace("&nbsp;", " ")
        .replace("&amp;", "&")
        .replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&#39;", "'")
    )
    # collapse whitespace
    text = re.sub(r"\s+", " ", text).strip()
    return text


def _maybe_clean_html_blob(text: str) -> str:
    """
    If the content looks like HTML (or contains a big HTML chunk),
    return a cleaner text version to improve retrieval.
    """
    if not text:
        return ""
    low = text.lstrip().lower()
    if low.startswith("<!doctype") or low.startswith("<html") or ("<body" in low and "</" in low):
        return _html_to_text(text)

    # If you store "[HTML]\n<...>", clean that part too
    if "\n[HTML]\n" in text:
        parts = text.split("\n[HTML]\n", 1)
        left = parts[0].strip()
        right = parts[1].strip()
        cleaned_right = _ = _html_to_text(right)
        out = left
        if cleaned_right:
            out = (out + "\n\n" + cleaned_right).strip() if out else cleaned_right
        return out

    return text


def _parse_pdf(data: bytes) -> str:
    if PdfReader is None:
        return ""
    try:
        reader = PdfReader(io.BytesIO(data))
        parts: List[str] = []
        for p in reader.pages:
            try:
                parts.append(p.extract_text() or "")
            except Exception:
                parts.append("")
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _parse_docx(data: bytes) -> str:
    if DocxDocument is None:
        return ""
    try:
        bio = io.BytesIO(data)
        doc = DocxDocument(bio)
        lines = []
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                lines.append(t)
        return "\n".join(lines).strip()
    except Exception:
        return ""


def _parse_xlsx(data: bytes) -> str:
    if openpyxl is None:
        return ""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
        out_lines: List[str] = []
        for sheet in wb.worksheets:
            out_lines.append(f"=== SHEET: {sheet.title} ===")
            for row in sheet.iter_rows(values_only=True):
                vals = []
                for v in row:
                    vals.append("" if v is None else str(v))
                line = ",".join(vals).strip()
                if line:
                    out_lines.append(line)
        return "\n".join(out_lines).strip()
    except Exception:
        return ""


def _parse_pptx(data: bytes) -> str:
    if Presentation is None:
        return ""
    try:
        prs = Presentation(io.BytesIO(data))
        parts: List[str] = []
        for i, slide in enumerate(prs.slides, start=1):
            parts.append(f"=== SLIDE {i} ===")
            for shape in slide.shapes:
                text = getattr(shape, "text", None)
                if text:
                    t = str(text).strip()
                    if t:
                        parts.append(t)
        return "\n".join(parts).strip()
    except Exception:
        return ""


def _parse_msg(data: bytes) -> str:
    if extract_msg is None:
        return ""
    try:
        with tempfile.NamedTemporaryFile(suffix=".msg", delete=True) as tmp:
            tmp.write(data)
            tmp.flush()
            msg = extract_msg.Message(tmp.name)
            msg.process()
            parts = []
            if msg.subject:
                parts.append(f"Subject: {msg.subject}")
            if msg.sender:
                parts.append(f"From: {msg.sender}")
            if msg.date:
                parts.append(f"Date: {msg.date}")
            body = (msg.body or "").strip()
            if body:
                parts.append(body)
            return "\n".join(parts).strip()
    except Exception:
        return ""


def _eml_get_text_and_html(message: Message) -> Tuple[str, str]:
    text_parts: List[str] = []
    html_parts: List[str] = []

    if message.is_multipart():
        for part in message.walk():
            ctype = (part.get_content_type() or "").lower()
            disp = (part.get_content_disposition() or "").lower()
            # skip attachments inside eml
            if disp == "attachment":
                continue
            try:
                payload = part.get_content()
            except Exception:
                payload = None

            if ctype == "text/plain" and isinstance(payload, str):
                t = payload.strip()
                if t:
                    text_parts.append(t)
            elif ctype == "text/html" and isinstance(payload, str):
                h = payload.strip()
                if h:
                    html_parts.append(h)
    else:
        ctype = (message.get_content_type() or "").lower()
        try:
            payload = message.get_content()
        except Exception:
            payload = None
        if ctype == "text/plain" and isinstance(payload, str):
            text_parts.append(payload.strip())
        elif ctype == "text/html" and isinstance(payload, str):
            html_parts.append(payload.strip())

    return ("\n\n".join(text_parts).strip(), "\n\n".join(html_parts).strip())


def _parse_eml(data: bytes) -> str:
    """
    Parse .eml reliably:
    - Extract headers (subject/from/date)
    - Extract body text/plain, and fallback to cleaned html
    """
    try:
        msg = BytesParser(policy=policy.default).parsebytes(data)
    except Exception:
        return ""

    subject = (msg.get("subject") or "").strip()
    from_ = (msg.get("from") or "").strip()
    date_ = (msg.get("date") or "").strip()

    text_body, html_body = _eml_get_text_and_html(msg)
    if not text_body and html_body:
        text_body = _html_to_text(html_body)

    parts: List[str] = []
    if subject:
        parts.append(f"Subject: {subject}")
    if from_:
        parts.append(f"From: {from_}")
    if date_:
        parts.append(f"Date: {date_}")
    if text_body:
        parts.append(text_body)

    return "\n".join(parts).strip()


def parse_file_bytes(filename: str, data: bytes) -> Tuple[bool, str, Optional[str]]:
    name = (filename or "").lower()
    ext = name.split(".")[-1] if "." in name else ""

    text = ""
    try:
        if ext == "pdf":
            text = _parse_pdf(data)
        elif ext in ("docx", "doc"):
            text = _parse_docx(data)
        elif ext in ("xlsx", "xls"):
            text = _parse_xlsx(data)
        elif ext in ("pptx", "ppt"):
            text = _parse_pptx(data)
        elif ext == "msg":
            text = _parse_msg(data)
        elif ext == "eml":
            text = _parse_eml(data)
        else:
            # Fallback decode as text
            try:
                text = data.decode("utf-8", errors="ignore").strip()
            except Exception:
                text = ""
    except Exception as e:
        return False, "", str(e)

    # Critical: clean HTML blobs for better retrieval
    text = _maybe_clean_html_blob(text)

    if not text:
        return True, "", None

    return True, text, None


def create_document(
    db: Session,
    *,
    name: str,
    source: str,
    source_kind: str,
    doc_date_value: Optional[date],
    content_text: str,
    email_message_id: Optional[str] = None,
) -> Document:
    doc = Document(
        name=name,
        source=source,
        source_kind=source_kind,
        doc_date=doc_date_value,
        content_text=content_text,
        email_message_id=email_message_id,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def _save_bytes_to_uploads(*, subdir: str, filename: str, content: bytes) -> Optional[str]:
    base_dir = os.getenv("UPLOADS_DIR", "/data/uploads")
    safe_subdir = re.sub(r"[^a-zA-Z0-9._/\-]", "_", subdir).strip("/")

    safe_name = os.path.basename(filename or "file.bin")
    safe_name = re.sub(r"[^a-zA-Z0-9._\-\s]", "_", safe_name)

    path_dir = os.path.join(base_dir, safe_subdir)
    path = os.path.join(path_dir, safe_name)

    try:
        os.makedirs(path_dir, exist_ok=True)
        with open(path, "wb") as f:
            f.write(content)
        return path
    except Exception:
        return None


# ============================
#  Public ingest helpers (Email Poller depends on these)
# ============================

def ingest_text_document(
    *,
    name: str,
    text: str,
    source: str,
    doc_date: str,
    email_message_id: Optional[str],
    source_kind: str,
) -> int:
    try:
        doc_date_value = date.fromisoformat(doc_date) if doc_date else date.today()
    except Exception:
        doc_date_value = date.today()

    # Important: clean HTML noise so it becomes searchable
    cleaned = _maybe_clean_html_blob(text or "")

    db = SessionLocal()
    try:
        doc = create_document(
            db,
            name=name,
            source=source,
            source_kind=source_kind,
            doc_date_value=doc_date_value,
            content_text=cleaned or "",
            email_message_id=email_message_id,
        )
        _safe_print(f'[INFO] 📧 Email Body Ingested: {name} | Source: {source}')
        return doc.id
    finally:
        db.close()


def ingest_uploaded_file_bytes(
    *,
    filename: str,
    content: bytes,
    source: str,
    doc_date: str,
    email_message_id: Optional[str],
    source_kind: str,
) -> int:
    try:
        doc_date_value = date.fromisoformat(doc_date) if doc_date else date.today()
    except Exception:
        doc_date_value = date.today()

    # Persist raw attachment for traceability
    subdir = "email_attachments"
    if email_message_id:
        subdir = f"email_attachments/{email_message_id}"
    _save_bytes_to_uploads(subdir=subdir, filename=filename, content=content)

    parsed_ok, text, parse_error = parse_file_bytes(filename, content)

    db = SessionLocal()
    try:
        doc = create_document(
            db,
            name=filename or "attachment",
            source=source,
            source_kind=source_kind,
            doc_date_value=doc_date_value,
            content_text=text or "",
            email_message_id=email_message_id,
        )
        if parse_error:
            _safe_print(f'[INFO] 📎 Attachment Ingested with parse error: {filename} | err: {parse_error}')
        else:
            _safe_print(f'[INFO] 📎 Attachment Ingested: {filename} | parsed={bool(parsed_ok)}')
        return doc.id
    finally:
        db.close()


# ============================
#  HTTP endpoint: upload
# ============================

@router.post("/upload")
async def upload_files(
    files: List[UploadFile] = File(...),
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    ingested = []
    for f in files:
        name = f.filename or "unknown"
        try:
            data = await f.read()
        except Exception as e:
            ingested.append(
                {
                    "name": name,
                    "id": None,
                    "source": "Upload",
                    "date": "",
                    "parsed": False,
                    "embedded": "disabled",
                    "error": f"failed to read file: {e}",
                }
            )
            continue

        doc_date_value = _guess_doc_date_from_filename(name) or date.today()
        parsed_ok, text, parse_error = parse_file_bytes(name, data)

        doc = create_document(
            db,
            name=name,
            source="Upload",
            source_kind="upload",
            doc_date_value=doc_date_value,
            content_text=text or "",
            email_message_id=None,
        )

        _safe_print(f'[INFO] 📄 File Ingested: {name} | Source: Upload')

        ingested.append(
            {
                "name": name,
                "id": doc.id,
                "source": "Upload",
                "date": doc_date_value.isoformat() if doc_date_value else "",
                "parsed": bool(parsed_ok),
                "embedded": "ok" if os.environ.get("OPENAI_API_KEY") else "disabled",
                "error": parse_error,
            }
        )

    return {"ok": True, "ingested": ingested}
